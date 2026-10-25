"""Turn publisher files into the small tables the models read.

Person columns (managers, named individuals, email addresses) are never written
out. Social care is limited to the adult services the England register covers.
"""

from __future__ import annotations

import csv
from io import StringIO
from pathlib import Path

CIW_KEEP = {
    "Care Home Service": {"Adults With Nursing", "Adults Without Nursing"},
    "Domiciliary Support Service": None,
    "Adult Placement Service": None,
}
CIW_COLUMNS = {
    "Service URN": "service_urn",
    "Service name": "name",
    "Known as": "known_as",
    "Service type": "service_type",
    "Service sub-type": "service_subtype",
    "Service address line 1": "address_1",
    "Service address line 2": "address_2",
    "Service address line 3": "address_3",
    "Service town/city": "town",
    "Service postcode": "postcode",
    "Local authority": "local_authority",
    "Provider URN": "provider_urn",
    "Provider name": "provider_name",
    "Provider type": "provider_type",
    "Company number": "company_number",
    "Charity number": "charity_number",
    "Provider registration date": "provider_registered_on",
}

CARE_HOME_SUBTYPES = {
    "Older People",
    "Learning Disabilities",
    "Mental Health Problems",
    "Physical and Sensory Impairment",
    "Alcohol and Drugs Rehabilitation",
    "Respite Care and Short Breaks",
    "Blood Borne Virus",
}
INSPECTORATE_KEEP = {
    "Care Home Service": CARE_HOME_SUBTYPES,
    "Support Service": {"Care at Home"},
    "Housing Support Service": None,
    "Nurse Agency": None,
    "Adult Placement Service": None,
}
INSPECTORATE_COLUMNS = {
    "CSNumber": "service_id",
    "CareService": "service_type",
    "Subtype": "service_subtype",
    "ServiceName": "name",
    "Address_line_1": "address_1",
    "Address_line_2": "address_2",
    "Address_line_3": "address_3",
    "Service_town": "town",
    "Service_Postcode": "postcode",
    "SP_number": "provider_id",
    "ServiceProvider": "provider_name",
    "DateReg": "registered_on",
    "Health_Board_Name": "health_board_name",
    "Council_Area_Name": "council_name",
}

RQIA_CATEGORIES = {
    "Nursing (NH)",
    "Residential (RC)",
    "Domiciliary Care Agency (DCA)",
    "Nursing Agency (NA)",
    "Independent Hospital (IH)",
    "Independent Hospital (IH) - Dental Treatment",
    "Independent Clinic (IC)",
    "Independent Medical Agency (IMA)",
    "Adult Placement Agency (APA)",
}
RQIA_FIELDS = {
    "Service ID": "service_id",
    "Service": "name",
    "Address Line1": "address_1",
    "Address Line2": "address_2",
    "Town": "town",
    "Postcode": "postcode",
    "Category": "category",
    "Provider": "provider_name",
}


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _clean(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\r", " ").replace("\n", " ").strip()
    return "" if text.lower() == "nan" else text


def _keep(service_type: str, subtype: str, allowed: dict[str, set[str] | None]) -> bool:
    rule = allowed.get(service_type)
    if rule is None and service_type not in allowed:
        return False
    if rule is None:
        return True
    return subtype in rule


def _write(dest: Path, fieldnames: list[str], rows: list[dict[str, str]]) -> int:
    dest.parent.mkdir(parents=True, exist_ok=True)
    with dest.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, lineterminator="\n", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


def _mapped_rows(source: Path, columns: dict[str, str], allowed: dict[str, set[str] | None]) -> list[dict[str, str]]:
    reader = csv.DictReader(StringIO(_read_text(source)))
    rows: list[dict[str, str]] = []
    for record in reader:
        service_type = _clean(record.get("Service type") or record.get("CareService"))
        subtype = _clean(record.get("Service sub-type") or record.get("Subtype"))
        if not _keep(service_type, subtype, allowed):
            continue
        mapped = {label: _clean(record.get(header)) for header, label in columns.items()}
        if mapped.get("service_urn") or mapped.get("service_id"):
            if mapped.get("name"):
                rows.append(mapped)
    return rows


def convert_ciw(source: Path, dest: Path) -> int:
    rows = _mapped_rows(source, CIW_COLUMNS, CIW_KEEP)
    return _write(dest, list(CIW_COLUMNS.values()), rows)


def convert_care_inspectorate(source: Path, dest: Path) -> int:
    statuses = {
        _clean(row.get("CSNumber")): _clean(row.get("ServiceStatus")).lower()
        for row in csv.DictReader(StringIO(_read_text(source)))
    }
    rows = [
        record
        for record in _mapped_rows(source, INSPECTORATE_COLUMNS, INSPECTORATE_KEEP)
        if statuses.get(record["service_id"], "active") in {"", "active"}
    ]
    return _write(dest, list(INSPECTORATE_COLUMNS.values()), rows)


def _rqia_header(rows: list[tuple]) -> int | None:
    for index, row in enumerate(rows[:15]):
        values = [_clean(cell) for cell in row]
        if "Service ID" in values and "Postcode" in values:
            return index
    return None


def convert_rqia(source: Path, dest: Path) -> int:
    import openpyxl

    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        grid = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    header_at = _rqia_header(grid)
    if header_at is None:
        raise RuntimeError("RQIA register has no Service ID header")
    headers = [_clean(cell) for cell in grid[header_at]]
    indexes = {label: headers.index(label) for label in RQIA_FIELDS if label in headers}
    missing = [label for label in RQIA_FIELDS if label not in indexes]
    if missing:
        raise RuntimeError("RQIA register is missing " + ", ".join(missing))
    rows: list[dict[str, str]] = []
    for raw in grid[header_at + 1:]:
        record = {
            label: _clean(raw[index] if index < len(raw) else "")
            for label, index in (
                (label, indexes[header]) for header, label in RQIA_FIELDS.items()
            )
        }
        if not record["service_id"] or not record["name"]:
            continue
        if record["category"] not in RQIA_CATEGORIES:
            continue
        rows.append(record)
    return _write(dest, list(RQIA_FIELDS.values()), rows)


def _csv_records(source: Path) -> list[dict[str, str]]:
    return [
        {name.strip(): _clean(value) for name, value in row.items() if name}
        for row in csv.DictReader(StringIO(_read_text(source)))
    ]


def convert_phs_hospitals(source: Path, dest: Path) -> int:
    rows = []
    for record in _csv_records(source):
        code = record.get("HospitalCode", "")
        name = record.get("HospitalName", "")
        if not code or not name:
            continue
        rows.append({
            "hospital_code": code,
            "name": name,
            "address_1": record.get("AddressLine1", ""),
            "address_2": record.get("AddressLine2", ""),
            "address_3": record.get("AddressLine3", ""),
            "address_4": record.get("AddressLine4", ""),
            "postcode": record.get("Postcode", ""),
            "health_board_code": record.get("HealthBoard", ""),
            "council_code": record.get("CouncilArea", ""),
        })
    return _write(dest, list(rows[0].keys()) if rows else [
        "hospital_code", "name", "address_1", "address_2", "address_3", "address_4",
        "postcode", "health_board_code", "council_code",
    ], rows)


def convert_phs_practices(source: Path, dest: Path) -> int:
    rows = []
    for record in _csv_records(source):
        code = record.get("PracticeCode", "")
        name = record.get("GPPracticeName", "")
        if not code or not name:
            continue
        rows.append({
            "practice_code": code,
            "name": name,
            "address_1": record.get("AddressLine1", ""),
            "address_2": record.get("AddressLine2", ""),
            "address_3": record.get("AddressLine3", ""),
            "address_4": record.get("AddressLine4", ""),
            "postcode": record.get("Postcode", ""),
            "health_board_code": record.get("HB", ""),
        })
    fields = [
        "practice_code", "name", "address_1", "address_2", "address_3", "address_4",
        "postcode", "health_board_code",
    ]
    return _write(dest, fields, rows)


def convert_phs_boards(source: Path, dest: Path) -> int:
    seen: set[str] = set()
    rows = []
    for record in _csv_records(source):
        if record.get("CADateArchived") or record.get("HBDateArchived"):
            continue
        code = record.get("CA", "")
        board = record.get("HB", "")
        name = record.get("HBName", "")
        if not code or not board or code in seen:
            continue
        seen.add(code)
        rows.append({
            "council_code": code,
            "council_name": record.get("CAName", ""),
            "health_board_code": board,
            "health_board_name": name,
        })
    return _write(
        dest,
        ["council_code", "council_name", "health_board_code", "health_board_name"],
        rows,
    )


CONVERTERS = {
    "ciw": convert_ciw,
    "care_inspectorate": convert_care_inspectorate,
    "rqia": convert_rqia,
    "phs_hospitals": convert_phs_hospitals,
    "phs_practices": convert_phs_practices,
    "phs_boards": convert_phs_boards,
}

OUTPUT_FIELDS = {
    "ciw": list(CIW_COLUMNS.values()),
    "care_inspectorate": list(INSPECTORATE_COLUMNS.values()),
    "rqia": list(RQIA_FIELDS.values()),
    "phs_hospitals": [
        "hospital_code", "name", "address_1", "address_2", "address_3", "address_4",
        "postcode", "health_board_code", "council_code",
    ],
    "phs_practices": [
        "practice_code", "name", "address_1", "address_2", "address_3", "address_4",
        "postcode", "health_board_code",
    ],
    "phs_boards": ["council_code", "council_name", "health_board_code", "health_board_name"],
}
